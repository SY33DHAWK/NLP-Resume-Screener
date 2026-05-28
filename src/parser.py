import pdfplumber
import docx
import os


def extract_text_from_pdf(file_path: str) -> str:
    """Extract raw text from a PDF file."""
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    return text.strip()


def extract_text_from_docx(file_path: str) -> str:
    """Extract raw text from a DOCX file."""
    doc = docx.Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    return text.strip()


def extract_text_from_txt(file_path: str) -> str:
    """Extract raw text from a TXT file."""
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read().strip()


def parse_document(file_path: str) -> str:
    """
    Master parser — detects file type and routes
    to the correct extractor automatically.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)
    elif extension == ".docx":
        return extract_text_from_docx(file_path)
    elif extension == ".txt":
        return extract_text_from_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")


def parse_from_uploaded_file(uploaded_file) -> str:
    """
    Handles files uploaded directly through Streamlit.
    Accepts Streamlit's UploadedFile object.
    """
    extension = os.path.splitext(uploaded_file.name)[1].lower()

    if extension == ".pdf":
        with pdfplumber.open(uploaded_file) as pdf:
            text = ""
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        return text.strip()

    elif extension == ".docx":
        doc = docx.Document(uploaded_file)
        text = ""
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
        return text.strip()

    elif extension == ".txt":
        return uploaded_file.read().decode("utf-8").strip()

    else:
        raise ValueError(f"Unsupported file type: {extension}")